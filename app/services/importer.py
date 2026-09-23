import csv
import io
from openpyxl import load_workbook
from app.models import Engineer, Job

def _parse_skills(value):
    # GPON:3;ETHERNET:2
    result = {}
    if not value:
        return result
    for part in str(value).split(";"):
        part = part.strip()
        if not part:
            continue
        if ":" in part:
            name, level = part.split(":", 1)
            result[name.strip().upper()] = int(level)
        else:
            result[part.upper()] = 1
    return result

def _parse_equipment(value):
    if not value:
        return []
    return [x.strip().upper() for x in str(value).split(";") if x.strip()]

def _to_minutes(value):
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    s = str(value).strip()
    if ":" in s:
        h, m = s.split(":", 1)
        return int(h) * 60 + int(m)
    return int(s)

def _rows_from_csv(raw):
    text = raw.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))

def _rows_from_xlsx(raw, sheet_name=None):
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    return [
        {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
        for row in rows[1:] if any(v is not None for v in row)
    ]

def parse_engineers(raw, filename):
    rows = _rows_from_xlsx(raw, "engineers") if filename.lower().endswith(".xlsx") else _rows_from_csv(raw)
    items = []
    for r in rows:
        items.append(Engineer(
            id=str(r["id"]).strip(),
            name=str(r.get("name") or r["id"]).strip(),
            lat=float(r["lat"]),
            lon=float(r["lon"]),
            shift_start=_to_minutes(r["shift_start"]),
            shift_end=_to_minutes(r["shift_end"]),
            skills=_parse_skills(r.get("skills")),
            equipment=_parse_equipment(r.get("equipment")),
        ))
    return items

def parse_jobs(raw, filename):
    rows = _rows_from_xlsx(raw, "jobs") if filename.lower().endswith(".xlsx") else _rows_from_csv(raw)
    items = []
    for r in rows:
        items.append(Job(
            id=str(r["id"]).strip(),
            lat=float(r["lat"]),
            lon=float(r["lon"]),
            priority=int(r["priority"]),
            service_minutes=int(r["service_minutes"]),
            window_start=_to_minutes(r["window_start"]),
            window_end=_to_minutes(r["window_end"]),
            sla_deadline=_to_minutes(r["sla_deadline"]),
            required_skills=_parse_skills(r.get("required_skills")),
            required_equipment=_parse_equipment(r.get("required_equipment")),
        ))
    return items
