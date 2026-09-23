import csv
import io
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from app.models import Engineer, Job
from app.services.beeline_data import (
    REGIONS,
    _engineer_attributes,
    _normative_for,
    _point_from_district,
    _priority,
    _required_skill,
)

DEFAULT_CENTER = (55.751244, 37.618423)


def _clean(value):
    return "" if value is None else str(value).strip()


def _first(row, *keys):
    for key in keys:
        value = row.get(key)
        if value not in (None, ""):
            return value
    return None


def _float(value, field, row_no, default=None):
    if value in (None, ""):
        if default is not None:
            return float(default)
        raise ValueError(f"строка {row_no}: отсутствует поле {field}")
    try:
        return float(str(value).replace(",", "."))
    except Exception:
        raise ValueError(f"строка {row_no}: поле {field} должно быть числом, получено {value!r}")


def _int(value, field, row_no, default=None):
    if value in (None, ""):
        if default is not None:
            return int(default)
        raise ValueError(f"строка {row_no}: отсутствует поле {field}")
    try:
        return int(float(str(value).replace(",", ".")))
    except Exception:
        raise ValueError(f"строка {row_no}: поле {field} должно быть целым числом, получено {value!r}")


def _minutes(value, default=None):
    if value in (None, ""):
        if default is not None:
            return int(default)
        raise ValueError("пустое значение времени")
    if isinstance(value, (int, float)):
        # Excel can store time as a fraction of day.
        if 0 <= value < 1:
            return int(round(value * 24 * 60))
        return int(value)
    text = str(value).strip()
    for fmt in ("%H:%M", "%d.%m.%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            dt = datetime.strptime(text, fmt)
            return dt.hour * 60 + dt.minute
        except ValueError:
            pass
    if text.isdigit():
        return int(text)
    raise ValueError(f"неизвестный формат времени: {text!r}")


def _rows(raw: bytes, filename: str):
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [_clean(x) for x in values[0]]
        return [
            {headers[i]: row[i] for i in range(len(headers))}
            for row in values[1:]
            if any(x not in (None, "") for x in row)
        ]
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251")
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=";,").delimiter
    except Exception:
        delimiter = ";"
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def _infer_region(filename: str, rows):
    hay = filename.lower()
    if "юго-вост" in hay or "юговост" in hay:
        return "southeast"
    if "югоцентр" in hay or "юго-центр" in hay:
        return "southcenter"
    if "восток" in hay:
        return "east"
    # Try explicit region column.
    for row in rows[:10]:
        region = _clean(_first(row, "region", "Регион")).lower()
        if "юго-вост" in region:
            return "southeast"
        if "югоцентр" in region:
            return "southcenter"
        if region == "восток":
            return "east"
    return None


def _office_for(region_key: Optional[str]):
    return REGIONS.get(region_key or "", {}).get("office", DEFAULT_CENTER)


def _looks_like_control(rows):
    return bool(rows and "Бригада" in rows[0])


def _looks_like_organizer_jobs(rows):
    if not rows:
        return False
    keys = set(rows[0].keys())
    return {"Заявка", "Тип заявки BK", "Начало", "Окончание", "Район", "Адрес"}.issubset(keys)


def parse_engineers(raw: bytes, filename: str):
    rows = _rows(raw, filename)
    if not rows:
        raise ValueError("файл инженеров пуст")

    region_key = _infer_region(filename, rows)
    office = _office_for(region_key)

    # Organizer control-distribution CSV can be imported as a brigade roster.
    if _looks_like_control(rows) and not any(k in rows[0] for k in ("lat", "latitude", "Широта")):
        names = []
        seen = set()
        for row in rows:
            name = _clean(row.get("Бригада"))
            if name and name not in seen:
                seen.add(name)
                names.append(name)
        if not names:
            raise ValueError("в колонке «Бригада» не найдено ни одной бригады")
        result = []
        for i, name in enumerate(names):
            skills, vehicle = _engineer_attributes(i)
            result.append(Engineer(
                id=f"E{i+1:02d}",
                name=name,
                lat=office[0],
                lon=office[1],
                shift_start=9 * 60,
                shift_end=23 * 60,
                skills=skills,
                vehicle_type=vehicle,
                region=REGIONS.get(region_key or "", {}).get("name", "Пользовательский набор"),
            ))
        return result

    result = []
    for i, r in enumerate(rows, start=2):
        skills_raw = _clean(_first(r, "skills", "Навыки", "skill", "Навык"))
        skills = [x.strip() for x in skills_raw.replace(",", "|").split("|") if x.strip()]
        if not skills:
            raise ValueError(f"строка {i}: отсутствуют навыки (skills/Навыки)")

        lat_raw = _first(r, "lat", "latitude", "Широта")
        lon_raw = _first(r, "lon", "longitude", "Долгота")
        # If coordinates are omitted but a recognized region is given, use the office/start point.
        lat = _float(lat_raw, "lat/Широта", i, default=office[0])
        lon = _float(lon_raw, "lon/Долгота", i, default=office[1])

        result.append(Engineer(
            id=_clean(_first(r, "id", "ID")) or f"E{i-1:02d}",
            name=_clean(_first(r, "name", "Имя", "Бригада")) or f"Инженер {i-1}",
            lat=lat,
            lon=lon,
            shift_start=_minutes(_first(r, "shift_start", "Начало смены"), 9 * 60),
            shift_end=_minutes(_first(r, "shift_end", "Конец смены"), 18 * 60),
            skills=skills,
            vehicle_type=_clean(_first(r, "vehicle_type", "Транспорт", "Тип транспорта")) or "Автомобиль",
            region=_clean(_first(r, "region", "Регион")) or REGIONS.get(region_key or "", {}).get("name", "Пользовательский набор"),
        ))
    return result


def _organizer_job(row, idx, region_key, office):
    job_id = _clean(row.get("Заявка"))
    if not job_id.isdigit():
        return None
    district = _clean(row.get("Район"))
    address = _clean(row.get("Адрес"))
    bk = _clean(row.get("Тип заявки BK"))
    hd = _clean(row.get("Тип заявки HD"))
    ws = _minutes(row.get("Начало"))
    we = _minutes(row.get("Окончание"))
    lat, lon = _point_from_district(district, address, office)
    norm = _normative_for(bk)
    service = norm["technical_minutes"] + norm["documents_minutes"]
    return Job(
        id=job_id,
        source_index=idx,
        region=REGIONS.get(region_key or "", {}).get("name", "Пользовательский набор"),
        address=address,
        district=district,
        lat=lat,
        lon=lon,
        bk_type=bk,
        hd_type=hd,
        priority=_priority(hd),
        service_minutes=service,
        normative_name=norm["name"],
        normative_road_minutes=norm["road_minutes"],
        technical_minutes=norm["technical_minutes"],
        documents_minutes=norm["documents_minutes"],
        base_normative_minutes=norm["base_normative_minutes"],
        window_start=ws,
        window_end=we,
        sla_deadline=we,
        required_skill=_required_skill(bk),
        required_vehicle_type=None,
    )


def parse_jobs(raw: bytes, filename: str):
    rows = _rows(raw, filename)
    if not rows:
        raise ValueError("файл заявок пуст")

    region_key = _infer_region(filename, rows)
    office = _office_for(region_key)

    # Native organizer format: no lat/lon required.
    if _looks_like_organizer_jobs(rows):
        result = []
        for idx, row in enumerate(rows):
            job = _organizer_job(row, idx, region_key, office)
            if job is not None:
                result.append(job)
        if not result:
            raise ValueError("в файле не найдено строк заявок с числовым ID в колонке «Заявка»")
        return result

    result = []
    for i, r in enumerate(rows, start=2):
        vehicle = _clean(_first(r, "required_vehicle_type", "Требуемый транспорт")) or None
        window_end = _minutes(_first(r, "window_end", "Конец окна"), 18 * 60)
        sla_raw = _first(r, "sla_deadline", "SLA")

        address = _clean(_first(r, "address", "Адрес"))
        district = _clean(_first(r, "district", "Район"))
        lat_raw = _first(r, "lat", "latitude", "Широта")
        lon_raw = _first(r, "lon", "longitude", "Долгота")
        if lat_raw in (None, "") or lon_raw in (None, ""):
            if address or district:
                lat, lon = _point_from_district(district, address, office)
            else:
                raise ValueError(f"строка {i}: нет координат и нет адреса/района для резервного размещения")
        else:
            lat = _float(lat_raw, "lat/Широта", i)
            lon = _float(lon_raw, "lon/Долгота", i)

        bk = _clean(_first(r, "bk_type", "Тип заявки BK")) or "Пользовательская"
        hd = _clean(_first(r, "hd_type", "Тип заявки HD")) or "Пользовательская"
        norm = _normative_for(bk)
        explicit_service = _first(r, "service_minutes", "Длительность")
        service = _int(explicit_service, "service_minutes/Длительность", i, default=norm["technical_minutes"] + norm["documents_minutes"])

        result.append(Job(
            id=_clean(_first(r, "id", "ID", "Заявка")) or f"J{i-1:03d}",
            source_index=i - 2,
            region=_clean(_first(r, "region", "Регион")) or REGIONS.get(region_key or "", {}).get("name", "Пользовательский набор"),
            address=address,
            district=district,
            lat=lat,
            lon=lon,
            bk_type=bk,
            hd_type=hd,
            priority=_int(_first(r, "priority", "Приоритет"), "priority/Приоритет", i, default=_priority(hd)),
            service_minutes=service,
            normative_name=_clean(_first(r, "normative_name", "Норматив")) or norm["name"],
            normative_road_minutes=_int(_first(r, "normative_road_minutes", "Норматив дороги"), "normative_road_minutes", i, default=norm["road_minutes"]),
            technical_minutes=_int(_first(r, "technical_minutes", "Технические работы"), "technical_minutes", i, default=norm["technical_minutes"]),
            documents_minutes=_int(_first(r, "documents_minutes", "Документы"), "documents_minutes", i, default=norm["documents_minutes"]),
            base_normative_minutes=_int(_first(r, "base_normative_minutes", "Базовый норматив"), "base_normative_minutes", i, default=norm["base_normative_minutes"]),
            window_start=_minutes(_first(r, "window_start", "Начало окна", "Начало"), 9 * 60),
            window_end=window_end,
            sla_deadline=_minutes(sla_raw) if sla_raw not in (None, "") else window_end,
            required_skill=_clean(_first(r, "required_skill", "Требуемый навык")) or _required_skill(bk),
            required_vehicle_type=vehicle,
        ))
    return result
